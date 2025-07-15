import openai # type: ignore
from docx import Document # type: ignore
from datetime import datetime



address = "142 Washington Ave, Garden City, NY 11530"

# Initialize the OpenAI API Key
client = openai.OpenAI(api_key='sk-rSNepa2p0yahWHyJ27CuZd2snzGme9R2hlNOpreUWDT3BlbkFJLVmbTqYzzOpIt-K5mexrZw5W37ziLI2jkNiz0NmmAA')

print('''Calling AI (County, Market, Submarket)...''')

# Generate County, Market, Submarket
county_market_submarket = client.chat.completions.create(
    model="ft:gpt-4o-2024-08-06:personal:county-market-submarket1:A5fgZu33",  
    messages=[
        {"role": "system", "content": "Provide the County, City, State, Market and Submarket from the given address."},
        {"role": "user", "content": address}
    ]
)
print("Complete")

"""GENERATE EXECTUIVE SUMMARY"""

property_summary = client.chat.completions.create(
    model="gpt-4o",  
    messages=[
        {
            "role": "system",
            "content": (
                "Your task is to craft a single, detailed property summary paragraph. "
                "Use non-biased language and keep it similar to this example: "
                "'The subject is located in Salem, in Utah County. It is part of the Provo-Orem MSA. "
                "The subject property is located in northern Utah within the official boundaries of Utah County. "
                "The county is situated directly south of Salt Lake County. This area is generally called the Provo/Orem metropolitan area and is approximately 45 miles south of metropolitan Salt Lake, which is the financial center for the Intermountain Region. "
                "This region encompasses all of Utah, southern Idaho, southwestern Wyoming, and eastern Nevada. "
                "Utah County is part of a four-county area that is commonly known as the Wasatch Front. Provo is the Utah County seat.' "
                "In your response, ensure that you generate one coherent paragraph without additional commentary."
            )
        },
        {"role": "user", "content": address}
    ]
)
print("Property Summary Complete")


location_summary = client.chat.completions.create(
    model="gpt-4o",
    messages=[
        {
            "role": "system",
            "content": (
                "Your task is to craft a single, detailed location summary paragraph. Use non-biased language and follow the style of this example: "
                "'The subject is located on the corner of Hwy 198 and Elk Ridge Drive with good access and exposure. A major thoroughfare in the area is Hwy 198 which partially fronts the subject. "
                "The location also offers very close proximity to Salem Pond, Salem High School, Salem Community Center, Salem City Recreation, with limited retail areas in close proximity. "
                "The subject is surrounded by vacant land and residential uses. Utah County is broken up into three sectors: North County (Lindon to Lehi), Central County (Provo/Orem), and South County (Springville to Payson). "
                "Central County accounts for a lot of the class B office buildings. The following is taken from Reis, it shows the market area with an arrow pointing to the subject.' "
                "Generate one coherent paragraph without additional commentary."
            )
        },
        {"role": "user", "content": address}
    ]
)
print("Location Summary Complete")




'''GENERATE REGIONAL'''
print('''Calling AI (Regional)...''')

# Generate Regional Introduction Paragraph
regional_introduction = client.chat.completions.create(
    model="ft:gpt-4o-2024-08-06:personal:ra-test4:A0gmxi9k",  
    messages=[
        {"role": "system", "content": "Your job is to write a detailed and factual regional analysis about any given property address in one short paragraph."},
        {"role": "user", "content": address}
    ]
)
print("Introduction Complete")

# DEMOGRAPHIC ANALYSIS

# Generate Population Analysis Paragraph
population_analysis = client.chat.completions.create(
    model="ft:gpt-4o-2024-08-06:personal:pa-test1:A0tS2X2O",  
    messages=[
        {"role": "system", "content": "Your purpose is to craft an accurate population analyses for any given location, highlighting key demographic details, growth trends, and factors that contribute to the area's appeal and community dynamics, all within a concise, one-paragraph format. "},
        {"role": "user", "content": address}
    ]
)
print("Population Analysis Complete")

# Generate Education Analysis 
education_analysis = client.chat.completions.create(
    model="ft:gpt-4o-2024-08-06:personal:education1:A4EobkhJ",
    messages=[
        {"role": "system", "content": "Your purpose is to craft an Education analysis for any given location, within a concise, one-paragraph format."},
        {"role": "user", "content": address}
    ]
)
print("Education Analysis Complete")

# HOUSEHOLD TRENDS

# Generate Employment Paragraph
employment_analysis = client.chat.completions.create(
    model="ft:gpt-4o-2024-08-06:personal:employment1:A4GKieMd",
    messages=[
        {"role": "system", "content": "Your purpose is to craft an accurate employment analysis for any given location, within a concise, one-paragraph format."},
        {"role": "user", "content": address}
    ]
)
print("Employment Complete")

# Generate Residential Analysis Summary
residential_analysis = client.chat.completions.create(
    model="ft:gpt-4o-2024-08-06:personal:residential-analysis1:A4K3sVLX",
    messages=[
        {"role": "system", "content": "Your goal is to generate a residential analysis that concisely describes a city’s housing market, demographics, and principal economic sectors, emphasizing clarity and relevance in a single paragraph."},
        {"role": "user", "content": address}
    ]
)
print("Summary Complete")
print("Regional Analysis Complete...")



economic_factors = client.chat.completions.create(
    model="gpt-4o",
    messages=[
        {
            "role": "system",
            "content": (
                "Your task is to craft a single, detailed paragraph summarizing economic factors. Use non-biased language and follow this example closely: "
                "'Salem is a suburb of Payson and Provo/Orem market area. Salem is still considered somewhat of a rural area but over the years has begun to be built out. "
                "A majority of resident’s commute to other cities within the metropolitan area for employment. The largest industries in the city include manufacturing, public administration, agricultural uses and retail trade. "
                "The local economy consists of commercial and industrial businesses on the main arterials. The city’s commercial area is on Hwy 198, featuring retail, office, residential, and financial services.' "
                "Generate one coherent paragraph without additional commentary."
            )
        },
        {"role": "user", "content": address}
    ]
)
print("Economic Factors Summary Complete")


community_services = client.chat.completions.create(
    model="gpt-4o",
    messages=[
        {
            "role": "system",
            "content": (
                "Your task is to craft a single, detailed paragraph summarizing community services available in the surrounding area. "
                "Use non-biased language and follow this example: "
                "'Community services and facilities are readily available in the surrounding area. These include public services such as fire stations, hospitals, police stations, and schools (all ages). GreatSchools.org is an on-line tool that rates every school on a scale of one to ten based on test scores. They also track parents rating of the school on a one to five scale.' "
                "Generate one coherent paragraph without additional commentary."
            )
        },
        {"role": "user", "content": address}
    ]
)
print("Community Services Summary Complete")






# Replace Word Doc
doc = Document(r"C:\Users\IAMBE\OneDrive\Desktop\Python\Market Analysis\MarketTemplate2.docx")

today = datetime.today()
date = today.strftime("%B %d, %Y")


# Define placeholders and their replacements
placeholders = {
    '{{date}}': date,
    '{{property_summary}}': property_summary.choices[0].message.content,
    '{{location_summary}}': location_summary.choices[0].message.content,
    '{{regional_intro}}': regional_introduction.choices[0].message.content,
    '{{population_analysis}}': population_analysis.choices[0].message.content,
    '{{education_analysis}}': education_analysis.choices[0].message.content,
    '{{employment_analysis}}': employment_analysis.choices[0].message.content,
    '{{regional_outro}}': residential_analysis.choices[0].message.content,
    '{{economic_factors}}': economic_factors.choices[0].message.content,
    '{{community_services}}': community_services.choices[0].message.content,
}

# Function to replace placeholders
def replace_placeholder(paragraph, placeholders):
    for key, value in placeholders.items():
        if key in paragraph.text:
            paragraph.text = paragraph.text.replace(key, value)

# Apply replacements
for paragraph in doc.paragraphs:
    replace_placeholder(paragraph, placeholders)

# Save the updated document
output_file_path = r"C:\Users\IAMBE\OneDrive\Desktop\MRKTA.docx"
doc.save(output_file_path)
print("Done")